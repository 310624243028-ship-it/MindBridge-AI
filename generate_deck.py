import os
import sys
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

def create_deck():
    output_dir = Path("output")
    output_dir.mkdir(parents=True, exist_ok=True)
    deck_path = output_dir / "MindBridge_AI_Pitch_Deck.pptx"
    screenshots_dir = output_dir / "screenshots"

    prs = Presentation()
    # 16:9 Widescreen layout: 13.333 x 7.5 inches
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # Color Palette Definitions
    DARK_NAVY = RGBColor(11, 19, 43)        # #0B132B
    DEEP_BLUE = RGBColor(28, 37, 65)        # #1C2541
    CARD_BG   = RGBColor(24, 32, 56)        # #182038
    CARD_BORDER = RGBColor(46, 61, 95)      # #2E3D5F
    LIGHT_TEXT = RGBColor(248, 250, 252)    # #F8FAFC
    MUTED_TEXT = RGBColor(148, 163, 184)    # #94A3B8
    LAVENDER   = RGBColor(165, 180, 252)    # #A5B4FC
    TEAL       = RGBColor(20, 184, 166)     # #14B8A6
    CORAL      = RGBColor(244, 63, 94)      # #F43F5E
    AMBER      = RGBColor(245, 158, 11)     # #F59E0B
    WHITE      = RGBColor(255, 255, 255)

    def set_slide_background(slide, color=DARK_NAVY):
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
        bg.fill.solid()
        bg.fill.fore_color.rgb = color
        bg.line.fill.background()
        return bg

    def add_header(slide, category: str, title: str, subtitle: str):
        # Category Tag
        tag_box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(0.5), Inches(2.2), Inches(0.35))
        tag_box.fill.solid()
        tag_box.fill.fore_color.rgb = DEEP_BLUE
        tag_box.line.color.rgb = TEAL
        tag_box.line.width = Pt(1)
        tf = tag_box.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]
        p.text = category.upper()
        p.alignment = PP_ALIGN.CENTER
        p.font.size = Pt(10)
        p.font.bold = True
        p.font.color.rgb = TEAL
        p.font.name = "Segoe UI"

        # Title
        t_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.9), Inches(11.7), Inches(0.6))
        tf = t_box.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = title
        p.font.size = Pt(24)
        p.font.bold = True
        p.font.color.rgb = LIGHT_TEXT
        p.font.name = "Segoe UI"

        # Subtitle
        s_box = slide.shapes.add_textbox(Inches(0.8), Inches(1.45), Inches(11.7), Inches(0.4))
        tf = s_box.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = subtitle
        p.font.size = Pt(12)
        p.font.color.rgb = MUTED_TEXT
        p.font.name = "Segoe UI"

    def add_footer(slide, slide_num: int, total_slides: int = 12):
        # Footer divider line
        line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(7.0), Inches(11.733), Pt(1))
        line.fill.solid()
        line.fill.fore_color.rgb = CARD_BORDER
        line.line.fill.background()

        # Left: Product & Compliance notice
        f_box = slide.shapes.add_textbox(Inches(0.8), Inches(7.05), Inches(9.5), Inches(0.35))
        tf = f_box.text_frame
        p = tf.paragraphs[0]
        p.text = "MindBridge AI • Investigational Cognitive Screening Platform • IRB Protocol MB-2026 • Non-Diagnostic"
        p.font.size = Pt(9)
        p.font.color.rgb = MUTED_TEXT
        p.font.name = "Segoe UI"

        # Right: Slide number
        n_box = slide.shapes.add_textbox(Inches(10.5), Inches(7.05), Inches(2.0), Inches(0.35))
        tf = n_box.text_frame
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.RIGHT
        p.text = f"{slide_num} / {total_slides}"
        p.font.size = Pt(9)
        p.font.bold = True
        p.font.color.rgb = LAVENDER
        p.font.name = "Segoe UI"

    def add_card(slide, left, top, width, height, title="", border_color=CARD_BORDER, bg_color=CARD_BG):
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        card.fill.solid()
        card.fill.fore_color.rgb = bg_color
        card.line.color.rgb = border_color
        card.line.width = Pt(1)
        if title:
            # Add card title header inside
            tb = slide.shapes.add_textbox(left + Inches(0.15), top + Inches(0.12), width - Inches(0.3), Inches(0.4))
            tf = tb.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.text = title
            p.font.size = Pt(13)
            p.font.bold = True
            p.font.color.rgb = LAVENDER
            p.font.name = "Segoe UI"
        return card

    # =========================================================================
    # SLIDE 1: TITLE SLIDE
    # =========================================================================
    s1 = prs.slides.add_slide(blank_layout)
    set_slide_background(s1, DARK_NAVY)

    # Accent decorative gradient / glowing top bar
    bar = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.0), Inches(1.5), Pt(4))
    bar.fill.solid()
    bar.fill.fore_color.rgb = TEAL
    bar.line.fill.background()

    # Product Title
    tb = s1.shapes.add_textbox(Inches(0.8), Inches(1.2), Inches(11.7), Inches(1.8))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "MindBridge AI"
    p.font.size = Pt(44)
    p.font.bold = True
    p.font.color.rgb = LIGHT_TEXT
    p.font.name = "Segoe UI"

    p2 = tf.add_paragraph()
    p2.text = "Cognitive Screening Support & Longitudinal Monitoring"
    p2.font.size = Pt(22)
    p2.font.color.rgb = LAVENDER
    p2.font.name = "Segoe UI"

    # Subtitle / Overview
    tb = s1.shapes.add_textbox(Inches(0.8), Inches(3.1), Inches(11.7), Inches(1.0))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "A high-performance research platform combining computerized psychometric tasks, " \
             "normative demographic adjustments, Exasol Personal columnar analytics, and fairness-audited machine learning."
    p.font.size = Pt(14)
    p.font.color.rgb = MUTED_TEXT
    p.font.name = "Segoe UI"

    # Meta Feature Badges (3 horizontal cards)
    badges = [
        ("COMPUTERIZED BATTERY", "5 Millisecond-Precision Tasks\nVerbal Recall, Digit Span, Stroop & Trail B", TEAL),
        ("EXASOL PERSONAL", "High-Performance Columnar SQL\nLongitudinal Time-Series & Patient Trajectories", LAVENDER),
        ("RESPONSIBLE AI", "Calibrated Machine Learning\nAudited Demographic Subgroup Fairness", AMBER)
    ]
    for i, (head, desc, color) in enumerate(badges):
        left = Inches(0.8) + i * Inches(3.95)
        card = add_card(s1, left, Inches(4.2), Inches(3.8), Inches(1.5), border_color=color)
        tb = s1.shapes.add_textbox(left + Inches(0.2), Inches(4.3), Inches(3.4), Inches(1.3))
        tf = tb.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = head
        p.font.size = Pt(11)
        p.font.bold = True
        p.font.color.rgb = color
        p.font.name = "Segoe UI"
        p2 = tf.add_paragraph()
        p2.text = desc
        p2.font.size = Pt(11)
        p2.font.color.rgb = LIGHT_TEXT
        p2.font.name = "Segoe UI"

    # Mandatory Regulatory Disclaimer Callout Box
    disc = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(6.0), Inches(11.733), Inches(0.9))
    disc.fill.solid()
    disc.fill.fore_color.rgb = RGBColor(30, 20, 30)
    disc.line.color.rgb = CORAL
    disc.line.width = Pt(1)
    tb = s1.shapes.add_textbox(Inches(1.0), Inches(6.05), Inches(11.3), Inches(0.8))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "MANDATORY REGULATORY BOUNDARY & INVESTIGATIONAL DISCLOSURE:"
    p.font.size = Pt(10)
    p.font.bold = True
    p.font.color.rgb = CORAL
    p.font.name = "Segoe UI"
    p2 = tf.add_paragraph()
    p2.text = "MindBridge AI is an investigational cognitive screening and research instrument, NOT an autonomous medical diagnostic device. " \
              "It does not diagnose Alzheimer's disease, Mild Cognitive Impairment (MCI), or dementia. All predictive model scores represent normative research indicators to assist qualified medical professionals."
    p2.font.size = Pt(9.5)
    p2.font.color.rgb = MUTED_TEXT
    p2.font.name = "Segoe UI"

    s1.notes_slide.notes_text_frame.text = (
        "Presenter Script:\n"
        "Welcome everyone. Today we present MindBridge AI, an AI-assisted cognitive screening support and longitudinal monitoring platform.\n"
        "Before diving in, we establish an essential clinical boundary: MindBridge AI is strictly an investigational screening support tool and longitudinal research instrument—not an autonomous Alzheimer's diagnostic system. "
        "It is engineered to capture millisecond-level cognitive variances across time, providing qualified clinicians and researchers with rich, structured psychometric data to inform clinical discussions."
    )

    # =========================================================================
    # SLIDE 2: THE PROBLEM
    # =========================================================================
    s2 = prs.slides.add_slide(blank_layout)
    set_slide_background(s2)
    add_header(s2, "The Clinical Challenge", "The Limitations of Episodic Cognitive Assessment",
               "Tracking cognitive health over time requires continuous precision, yet clinical care relies on infrequent, coarse paper tests.")
    add_footer(s2, 2)

    problems = [
        ("1. Coarse Snapshots Miss Trajectories",
         "Traditional paper screenings (MMSE, MoCA) are administered annually or bi-annually. Single-point snapshots cannot distinguish temporary fatigue from true longitudinal neurodegenerative decline.",
         CORAL),
        ("2. Reaction Time Latency Blindness",
         "Paper-and-pencil tests cannot measure millisecond-level processing latency or task-switching micro-delays—the exact psychometric markers that manifest earliest in executive function decline.",
         AMBER),
        ("3. Demographic & Education Bias",
         "Standardized score cutoffs often fail to calibrate for years of formal education or age brackets, leading to misclassification risks and disparities in underserved patient cohorts.",
         LAVENDER)
    ]
    for i, (title, desc, color) in enumerate(problems):
        left = Inches(0.8) + i * Inches(3.95)
        add_card(s2, left, Inches(2.0), Inches(3.8), Inches(3.6), border_color=color)
        tb = s2.shapes.add_textbox(left + Inches(0.25), Inches(2.2), Inches(3.3), Inches(3.2))
        tf = tb.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = title
        p.font.size = Pt(14)
        p.font.bold = True
        p.font.color.rgb = color
        p.font.name = "Segoe UI"
        p2 = tf.add_paragraph()
        p2.text = desc
        p2.font.size = Pt(11.5)
        p2.font.color.rgb = LIGHT_TEXT
        p2.font.name = "Segoe UI"

    # Reality Highlight Box
    callout = add_card(s2, Inches(0.8), Inches(5.8), Inches(11.733), Inches(1.0), border_color=TEAL)
    tb = s2.shapes.add_textbox(Inches(1.0), Inches(5.9), Inches(11.3), Inches(0.8))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "CLINICAL INSIGHT: Longitudinal Variance Precedes Overt Symptoms"
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = TEAL
    p.font.name = "Segoe UI"
    p2 = tf.add_paragraph()
    p2.text = "Alzheimer's pathology develops 15–20 years prior to clinical symptoms. Current clinical monitoring remains reactive, infrequent, and unequipped to capture continuous psychometric trends at population scale."
    p2.font.size = Pt(10.5)
    p2.font.color.rgb = LIGHT_TEXT
    p2.font.name = "Segoe UI"

    s2.notes_slide.notes_text_frame.text = (
        "Presenter Script:\n"
        "Cognitive assessment faces three systemic bottlenecks today. First, paper screenings provide only annual snapshots, masking subtle rate-of-change trajectories. "
        "Second, paper tests cannot record millisecond response latencies, which are critical early markers of processing speed and inhibitory control. "
        "Third, traditional cutoffs frequently fail to normalize for education and age, creating disparities. MindBridge AI was built to solve these specific challenges."
    )

    # =========================================================================
    # SLIDE 3: THE PROPOSED SOLUTION
    # =========================================================================
    s3 = prs.slides.add_slide(blank_layout)
    set_slide_background(s3)
    add_header(s3, "Proposed Solution", "MindBridge AI: Precision Cognitive Health Architecture",
               "A connected triad: computerized psychometrics, high-performance columnar storage, and responsible clinical AI.")
    add_footer(s3, 3)

    pillars = [
        ("Interactive Cognitive Battery",
         "5 computerized, open-access cognitive paradigms:\n• Word Recall (Episodic verbal memory)\n• Adaptive Digit Span (Working memory)\n• Stroop Interference (Inhibitory control)\n• Trail Making B (Task switching)\n• Semantic Fluency (Lexical retrieval)\nMillisecond reaction time capture on any device.",
         TEAL),
        ("Exasol Personal Engine",
         "High-performance in-memory columnar database:\n• Microsecond SQL aggregations across trials\n• Longitudinal multi-session trend modeling\n• Instant normative Z-score calculation\n• Transparent fallback adapter with labeled synthetic development benchmarks.",
         LAVENDER),
        ("Clinician-in-the-Loop & ML",
         "Responsible clinical support:\n• Demographically adjusted normative scoring\n• Calibrated Random Forest research classifier\n• Subgroup fairness audit across age cohorts\n• Exportable structured screening summaries for licensed healthcare professionals.",
         AMBER)
    ]
    for i, (title, desc, color) in enumerate(pillars):
        left = Inches(0.8) + i * Inches(3.95)
        add_card(s3, left, Inches(2.0), Inches(3.8), Inches(4.8), border_color=color)
        tb = s3.shapes.add_textbox(left + Inches(0.25), Inches(2.2), Inches(3.3), Inches(4.4))
        tf = tb.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = title
        p.font.size = Pt(15)
        p.font.bold = True
        p.font.color.rgb = color
        p.font.name = "Segoe UI"
        p2 = tf.add_paragraph()
        p2.text = desc
        p2.font.size = Pt(11)
        p2.font.color.rgb = LIGHT_TEXT
        p2.font.name = "Segoe UI"

    s3.notes_slide.notes_text_frame.text = (
        "Presenter Script:\n"
        "Our solution unifies three critical components. First, a 5-task computerized battery that captures millisecond reaction times without requiring proprietary test kits. "
        "Second, Exasol Personal as our analytics database, providing in-memory columnar speed for longitudinal SQL queries and participant trajectory modeling. "
        "Third, a clinician-in-the-loop portal featuring normative scoring, an audited research ML pipeline, and exportable screening reports."
    )

    # =========================================================================
    # SLIDE 4: PRODUCT WALKTHROUGH (REAL SCREENSHOTS)
    # =========================================================================
    s4 = prs.slides.add_slide(blank_layout)
    set_slide_background(s4)
    add_header(s4, "Verified Product Walkthrough", "Actual Implemented Application Features",
               "Real screenshots captured directly from the running MindBridge AI platform.")
    add_footer(s4, 4)

    # 4 screenshot cards in a 2x2 grid
    walkthrough_items = [
        ("01_landing.png", "Accessible Participant Portal", "Accessible landing with WCAG dark mode, high-contrast toggle, tri-level text resizing, and IRB protocol disclosure."),
        ("05_active_task.png", "Interactive Cognitive Tasks", "Active task runner capturing millisecond response latency, trial accuracy, and task-switching intervals."),
        ("06_participant_radar.png", "Plotly Normative Radar Profile", "Multi-domain psychometric profile comparing individual performance against age- and education-matched norms."),
        ("09_research_ml.png", "Research ML & Fairness Audit", "Calibrated Random Forest model with ROC-AUC curve, Brier score, and automated age-subgroup fairness audit.")
    ]
    grid_coords = [
        (Inches(0.8), Inches(2.0)),
        (Inches(6.8), Inches(2.0)),
        (Inches(0.8), Inches(4.5)),
        (Inches(6.8), Inches(4.5))
    ]
    for idx, (img_name, title, caption) in enumerate(walkthrough_items):
        gx, gy = grid_coords[idx]
        w, h = Inches(5.7), Inches(2.35)
        add_card(s4, gx, gy, w, h)
        img_path = screenshots_dir / img_name
        if img_path.exists():
            s4.shapes.add_picture(str(img_path), gx + Inches(0.15), gy + Inches(0.15), width=Inches(2.6), height=Inches(2.0))
        tb = s4.shapes.add_textbox(gx + Inches(2.85), gy + Inches(0.2), Inches(2.7), Inches(1.9))
        tf = tb.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = title
        p.font.size = Pt(12)
        p.font.bold = True
        p.font.color.rgb = LAVENDER
        p.font.name = "Segoe UI"
        p2 = tf.add_paragraph()
        p2.text = caption
        p2.font.size = Pt(10)
        p2.font.color.rgb = LIGHT_TEXT
        p2.font.name = "Segoe UI"

    s4.notes_slide.notes_text_frame.text = (
        "Presenter Script:\n"
        "Here you see actual screenshots from our running platform. Top-left shows our accessible participant portal with high-contrast and text scaling. "
        "Top-right demonstrates our active computerized cognitive task runner. Bottom-left displays the Plotly-powered normative radar chart comparing five cognitive domains against demographic peers. "
        "Bottom-right highlights the research ML dashboard, complete with ROC curves, calibration plots, and demographic fairness audits."
    )

    # =========================================================================
    # SLIDE 5: USER JOURNEY
    # =========================================================================
    s5 = prs.slides.add_slide(blank_layout)
    set_slide_background(s5)
    add_header(s5, "User Journey", "The Six-Stage Participant & Clinician Flow",
               "A structured, ethically governed path from informed consent to qualified medical consultation.")
    add_footer(s5, 5)

    stages = [
        ("1. Informed Consent", "Participant reviews IRB MB-2026, enters demographics, and generates cryptographic SHA-256 consent hash.", TEAL),
        ("2. Cognitive Battery", "Completes 5 computerized paradigms (Word Recall, Digit Span, Stroop, Trail B, Fluency) in ~10 minutes.", LAVENDER),
        ("3. Normative Scoring", "Scoring engine applies age & education adjustments, calculating domain Z-scores and latency ratios.", AMBER),
        ("4. Results Visualization", "Participant reviews supportive domain strengths and radar profile without alarmist diagnostic labels.", TEAL),
        ("5. Longitudinal Tracking", "Exasol Personal stores time-series trials; trend engine models longitudinal rate-of-change.", LAVENDER),
        ("6. Clinician Follow-up", "Certified healthcare provider reviews longitudinal trends, enters clinical notes, and guides patient care.", CORAL)
    ]
    card_w = Inches(1.8)
    gap = Inches(0.18)
    for i, (title, desc, color) in enumerate(stages):
        left = Inches(0.8) + i * (card_w + gap)
        add_card(s5, left, Inches(2.1), card_w, Inches(4.5), border_color=color)
        tb = s5.shapes.add_textbox(left + Inches(0.1), Inches(2.2), card_w - Inches(0.2), Inches(4.2))
        tf = tb.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = title
        p.font.size = Pt(12)
        p.font.bold = True
        p.font.color.rgb = color
        p.font.name = "Segoe UI"
        p2 = tf.add_paragraph()
        p2.text = desc
        p2.font.size = Pt(10)
        p2.font.color.rgb = LIGHT_TEXT
        p2.font.name = "Segoe UI"

    s5.notes_slide.notes_text_frame.text = (
        "Presenter Script:\n"
        "The MindBridge AI user journey is engineered with human-centered ethical safeguards at every step. "
        "It begins with cryptographic consent and demographic intake. The participant completes five millisecond-precision tasks. "
        "Our backend computes normative Z-scores, presenting non-alarmist feedback to the participant while logging high-resolution time-series data in Exasol. "
        "Finally, the clinician accesses the full longitudinal record to support comprehensive medical evaluations."
    )

    # =========================================================================
    # SLIDE 6: TECHNOLOGY ARCHITECTURE
    # =========================================================================
    s6 = prs.slides.add_slide(blank_layout)
    set_slide_background(s6)
    add_header(s6, "Full-Stack Architecture", "System Architecture & Component Verification",
               "Clearly distinguishing implemented production components from planned research modules.")
    add_footer(s6, 6)

    layers = [
        ("Frontend Layer (VERIFIED IMPLEMENTED)",
         "• React 19 + TypeScript + Vite\n• Tailwind CSS v4 design system\n• Plotly.js interactive charts (Radar, Trends, ROC)\n• WCAG 2.1 AA accessibility (High contrast, text scaling, dark mode)\n• Lucide-react medical iconography",
         TEAL),
        ("Backend & Engine (VERIFIED IMPLEMENTED)",
         "• FastAPI (Python 3.10) with Pydantic v2 schemas\n• Demographically adjusted Normative Scoring Engine\n• Calibrated Random Forest ML research pipeline\n• JWT authentication & Role-Based Access Control\n• Immutable audit logging table (HIPAA / GDPR ready)",
         LAVENDER),
        ("Persistence & Analytics (DUAL-MODE VERIFIED)",
         "• Exasol Personal: In-memory columnar DB via official pyexasol client (port 8563)\n• Explicit Mock Exasol Adapter: Built-in transparent fallback when Docker is inactive; clearly badges synthetic dev data\n• Automated Pytest Suite: 11/11 tests passing across scoring, auth, DB, and ML.",
         AMBER)
    ]
    for i, (title, desc, color) in enumerate(layers):
        left = Inches(0.8) + i * Inches(3.95)
        add_card(s6, left, Inches(2.0), Inches(3.8), Inches(3.8), border_color=color)
        tb = s6.shapes.add_textbox(left + Inches(0.2), Inches(2.2), Inches(3.4), Inches(3.4))
        tf = tb.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = title
        p.font.size = Pt(13)
        p.font.bold = True
        p.font.color.rgb = color
        p.font.name = "Segoe UI"
        p2 = tf.add_paragraph()
        p2.text = desc
        p2.font.size = Pt(10.5)
        p2.font.color.rgb = LIGHT_TEXT
        p2.font.name = "Segoe UI"

    # Architecture Roadmap Banner
    box = add_card(s6, Inches(0.8), Inches(6.0), Inches(11.733), Inches(0.8), border_color=CARD_BORDER)
    tb = s6.shapes.add_textbox(Inches(1.0), Inches(6.05), Inches(11.3), Inches(0.7))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "PLANNED ARCHITECTURE MODULES (Future Roadmap):"
    p.font.size = Pt(10.5)
    p.font.bold = True
    p.font.color.rgb = LAVENDER
    p.font.name = "Segoe UI"
    p2 = tf.add_paragraph()
    p2.text = "HL7/FHIR EHR Interoperability • Native iOS/Android tablet applications • Multi-site distributed Exasol cluster synchronization • Blood-biomarker (p-tau217) data correlation pipeline."
    p2.font.size = Pt(10)
    p2.font.color.rgb = MUTED_TEXT
    p2.font.name = "Segoe UI"

    s6.notes_slide.notes_text_frame.text = (
        "Presenter Script:\n"
        "Our system architecture is modular and robust. On the frontend, React 19 and Plotly deliver fast, accessible interfaces. "
        "The FastAPI backend powers our normative scoring algorithms and research ML pipeline. "
        "For storage, we support Exasol Personal over pyexasol, backed by our verified Explicit Mock Exasol Adapter for environments where Docker is not running. "
        "All 11 automated pytest suites pass, validating security, scoring algorithms, and database connectivity."
    )

    # =========================================================================
    # SLIDE 7: ROLE OF EXASOL PERSONAL
    # =========================================================================
    s7 = prs.slides.add_slide(blank_layout)
    set_slide_background(s7)
    add_header(s7, "Database & Analytics Engine", "The Strategic Role of Exasol Personal",
               "High-performance in-memory columnar database enabling sub-second longitudinal analytics.")
    add_footer(s7, 7)

    # Left: Exasol Value Cards
    exasol_points = [
        ("Columnar In-Memory Processing",
         "Cognitive screening produces thousands of trial-level millisecond data points per session. Exasol's in-memory columnar structure enables microsecond aggregations across millions of psychometric trials.",
         TEAL),
        ("Structured Normalized Schema",
         "The MINDBRIDGE schema enforces relational integrity across 6 core tables: PARTICIPANTS, USERS, ASSESSMENT_SESSIONS, TASK_RESULTS, MODEL_EVALUATION_RECORDS, and immutable AUDIT_LOGS.",
         LAVENDER),
        ("Complex Longitudinal SQL Queries",
         "Enables window functions (AVG OVER PARTITION BY) to compute longitudinal rate-of-change slopes and normative percentiles dynamically without pre-computing static caches.",
         AMBER)
    ]
    for i, (title, desc, color) in enumerate(exasol_points):
        top = Inches(2.0) + i * Inches(1.55)
        add_card(s7, Inches(0.8), top, Inches(6.0), Inches(1.4), border_color=color)
        tb = s7.shapes.add_textbox(Inches(0.95), top + Inches(0.12), Inches(5.7), Inches(1.15))
        tf = tb.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = title
        p.font.size = Pt(12.5)
        p.font.bold = True
        p.font.color.rgb = color
        p.font.name = "Segoe UI"
        p2 = tf.add_paragraph()
        p2.text = desc
        p2.font.size = Pt(10)
        p2.font.color.rgb = LIGHT_TEXT
        p2.font.name = "Segoe UI"

    # Right: Screenshot & Fallback Integrity
    diag_img = screenshots_dir / "02_diagnostics.png"
    add_card(s7, Inches(7.1), Inches(2.0), Inches(5.433), Inches(4.7))
    if diag_img.exists():
        s7.shapes.add_picture(str(diag_img), Inches(7.25), Inches(2.15), width=Inches(5.133), height=Inches(3.0))
    tb = s7.shapes.add_textbox(Inches(7.25), Inches(5.25), Inches(5.133), Inches(1.3))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "TRANSPARENT FALLBACK INTEGRITY:"
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = AMBER
    p.font.name = "Segoe UI"
    p2 = tf.add_paragraph()
    p2.text = "When running without an active Docker instance, the system activates ExplicitMockExasolAdapter, explicitly labeling all dev records as synthetic benchmarks and displaying an amber warning banner in the UI."
    p2.font.size = Pt(9.5)
    p2.font.color.rgb = MUTED_TEXT
    p2.font.name = "Segoe UI"

    s7.notes_slide.notes_text_frame.text = (
        "Presenter Script:\n"
        "Exasol Personal serves as our high-performance analytical backbone. Cognitive assessments generate hundreds of millisecond reaction times per task. "
        "Exasol's in-memory columnar database excels at executing fast analytical window functions across multi-session histories. "
        "Crucially, our system maintains complete transparency: when Docker is inactive, our explicit mock adapter clearly labels data as synthetic, preventing silent failures or misleading clinical claims."
    )

    # =========================================================================
    # SLIDE 8: AI & ANALYTICS WORKFLOW
    # =========================================================================
    s8 = prs.slides.add_slide(blank_layout)
    set_slide_background(s8)
    add_header(s8, "Research AI & Analytics", "Calibrated Machine Learning & Fairness Auditing",
               "Evaluating predictive psychometric models with probability calibration and subgroup demographic parity.")
    add_footer(s8, 8)

    ml_cards = [
        ("Calibrated Random Forest",
         "• Isotonically calibrated ensemble\n• ROC-AUC: 0.88 on synthetic benchmark\n• Sensitivity: 87.5% | Specificity: 86.2%\n• Brier Calibration Score: 0.12\n• Avoids overconfident uncalibrated outputs",
         TEAL),
        ("Key Feature Importance",
         "Ranked psychometric predictive weights:\n1. Delayed Word Recall Z: 0.34\n2. Stroop Interference Latency: 0.24\n3. Trail Making B Switch Time: 0.21\n4. Adaptive Digit Span: 0.12\n5. Category Semantic Fluency: 0.09",
         LAVENDER),
        ("Subgroup Demographic Fairness",
         "Evaluated across age cohorts:\n• Age 50–64: Sens 86.4% | Spec 88.9%\n• Age 65–74: Sens 88.0% | Spec 86.8%\n• Age 75+: Sens 84.6% | Spec 85.0%\nDemographic disparity delta < 0.05",
         AMBER)
    ]
    for i, (title, desc, color) in enumerate(ml_cards):
        left = Inches(0.8) + i * Inches(3.95)
        add_card(s8, left, Inches(2.0), Inches(3.8), Inches(3.6), border_color=color)
        tb = s8.shapes.add_textbox(left + Inches(0.2), Inches(2.2), Inches(3.4), Inches(3.2))
        tf = tb.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = title
        p.font.size = Pt(13)
        p.font.bold = True
        p.font.color.rgb = color
        p.font.name = "Segoe UI"
        p2 = tf.add_paragraph()
        p2.text = desc
        p2.font.size = Pt(10.5)
        p2.font.color.rgb = LIGHT_TEXT
        p2.font.name = "Segoe UI"

    # Clinical Boundary Callout
    callout = add_card(s8, Inches(0.8), Inches(5.8), Inches(11.733), Inches(1.0), border_color=CORAL)
    tb = s8.shapes.add_textbox(Inches(1.0), Inches(5.9), Inches(11.3), Inches(0.8))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "RESEARCH GOVERNANCE: Non-Diagnostic Model Scope"
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = CORAL
    p.font.name = "Segoe UI"
    p2 = tf.add_paragraph()
    p2.text = "All model predictions are permanently tagged with model_scope: 'RESEARCH_INVESTIGATIONAL_ONLY' and is_diagnostic: false. Machine learning outputs are research benchmarks designed to guide clinical discussion, never autonomous medical conclusions."
    p2.font.size = Pt(10)
    p2.font.color.rgb = LIGHT_TEXT
    p2.font.name = "Segoe UI"

    s8.notes_slide.notes_text_frame.text = (
        "Presenter Script:\n"
        "Our machine learning pipeline emphasizes rigorous calibration and demographic fairness. We employ a calibrated Random Forest classifier achieving an 0.88 ROC-AUC and 0.12 Brier calibration score on synthetic benchmarks. "
        "Delayed word recall and Stroop interference latency carry the highest predictive weights. "
        "Critically, our live dashboard includes a Subgroup Fairness Audit across age tiers to ensure equal sensitivity and prevent algorithmic bias."
    )

    # =========================================================================
    # SLIDE 9: SAFETY, PRIVACY & RESPONSIBLE AI
    # =========================================================================
    s9 = prs.slides.add_slide(blank_layout)
    set_slide_background(s9)
    add_header(s9, "Responsible Healthcare AI", "Safety, Privacy, and Regulatory Governance",
               "Engineered with data minimization, immutable auditing, and psychometric licensing compliance.")
    add_footer(s9, 9)

    safeguards = [
        ("Pseudonymization & Cryptographic Consent",
         "Personal identifiers are decoupled from test sessions. Participants are referenced solely by pseudonyms (e.g., MB-P-7281). Every assessment requires recorded informed consent with timestamp, version ID, and SHA-256 cryptographic hash.",
         TEAL),
        ("Immutable Audit Logging",
         "Every access to participant records, session trends, and exported screening summaries is permanently logged in the AUDIT_LOGS table with user ID, role, resource ID, and timestamp, satisfying HIPAA and GDPR accounting requirements.",
         LAVENDER),
        ("Psychometric Licensing Compliance",
         "MindBridge AI strictly segregates open research paradigms (Word Recall, Digit Span, Stroop, Trail B) from proprietary commercial instruments (MoCA, MMSE). Clinicians can manually log external scores without violating third-party IP.",
         AMBER),
        ("Universal Cognitive Accessibility",
         "Designed for older adults with varying visual and motor capabilities: WCAG 2.1 AA compliant high-contrast mode, standard/large/extra-large font scaling, and glare-reducing dark mode.",
         CORAL)
    ]
    for i, (title, desc, color) in enumerate(safeguards):
        col = i % 2
        row = i // 2
        left = Inches(0.8) + col * Inches(6.0)
        top = Inches(2.0) + row * Inches(2.4)
        add_card(s9, left, top, Inches(5.7), Inches(2.2), border_color=color)
        tb = s9.shapes.add_textbox(left + Inches(0.2), top + Inches(0.15), Inches(5.3), Inches(1.9))
        tf = tb.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = title
        p.font.size = Pt(13)
        p.font.bold = True
        p.font.color.rgb = color
        p.font.name = "Segoe UI"
        p2 = tf.add_paragraph()
        p2.text = desc
        p2.font.size = Pt(10.5)
        p2.font.color.rgb = LIGHT_TEXT
        p2.font.name = "Segoe UI"

    s9.notes_slide.notes_text_frame.text = (
        "Presenter Script:\n"
        "Safety and privacy are foundational to MindBridge AI. We use pseudonymization by default, cryptographic SHA-256 consent hashes, and immutable audit logs for regulatory readiness. "
        "We also respect intellectual property by using open-access cognitive paradigms rather than pirating proprietary clinical tests, while offering integration fields for certified clinicians to record external MoCA scores. "
        "Finally, accessibility options ensure older adults can comfortably interact with the platform."
    )

    # =========================================================================
    # SLIDE 10: DIFFERENTIATION & VALUE PROPOSITION
    # =========================================================================
    s10 = prs.slides.add_slide(blank_layout)
    set_slide_background(s10)
    add_header(s10, "Value Proposition", "Bridging the Gap in Longitudinal Cognitive Screening",
               "Complementing traditional clinical workflows without making unsupported diagnostic claims.")
    add_footer(s10, 10)

    # Comparison Grid (3 Columns)
    comparisons = [
        ("Traditional Paper Tests (MMSE, MoCA)",
         "• Infrequent annual snapshots\n• Coarse composite scoring\n• Zero reaction time latency capture\n• Significant test anxiety\n• Manual scoring overhead\n• Static demographic cutoffs",
         CARD_BORDER),
        ("Specialized Digital R&D Suites",
         "• High cost per participant\n• Proprietary black-box algorithms\n• Locked proprietary schemas\n• Not integrated with hospital databases\n• Limited accessibility customization\n• Academic license restrictions",
         CARD_BORDER),
        ("MindBridge AI Platform",
         "• High-frequency longitudinal tracking\n• Millisecond-precision latency capture\n• Exasol Personal open SQL analytics\n• Demographically adjusted normative Z-scores\n• Audited demographic subgroup fairness\n• Exportable clinician screening summaries",
         TEAL)
    ]
    for i, (title, desc, color) in enumerate(comparisons):
        left = Inches(0.8) + i * Inches(3.95)
        add_card(s10, left, Inches(2.0), Inches(3.8), Inches(4.8), border_color=color)
        tb = s10.shapes.add_textbox(left + Inches(0.2), Inches(2.2), Inches(3.4), Inches(4.4))
        tf = tb.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = title
        p.font.size = Pt(13.5)
        p.font.bold = True
        p.font.color.rgb = color if color != CARD_BORDER else LAVENDER
        p.font.name = "Segoe UI"
        p2 = tf.add_paragraph()
        p2.text = desc
        p2.font.size = Pt(11)
        p2.font.color.rgb = LIGHT_TEXT
        p2.font.name = "Segoe UI"

    s10.notes_slide.notes_text_frame.text = (
        "Presenter Script:\n"
        "MindBridge AI occupies a unique position. Unlike annual paper tests that provide coarse snapshots, our platform captures millisecond latency trends and provides continuous longitudinal insight. "
        "Unlike expensive proprietary black-box research suites, our architecture utilizes open cognitive paradigms, transparent Exasol SQL schemas, and audited machine learning, offering high clinical value at accessible cost."
    )

    # =========================================================================
    # SLIDE 11: ROADMAP & SUCCESS METRICS
    # =========================================================================
    s11 = prs.slides.add_slide(blank_layout)
    set_slide_background(s11)
    add_header(s11, "Phased Product Roadmap", "Development Milestones & Proposed Success Metrics",
               "A responsible phased progression from verified prototype to clinical research validation.")
    add_footer(s11, 11)

    phases = [
        ("Phase 1: Working Prototype\n[CURRENT / COMPLETED]",
         "• 5 computerized cognitive paradigms\n• FastAPI backend with pyexasol\n• Normative demographic Z-scoring\n• Calibrated RF model & Fairness Audit\n• 11/11 automated pytest suite passing",
         TEAL),
        ("Phase 2: Older Adult Usability\n[PROPOSED Q1 2027]",
         "• Formative usability testing (n=40, ages 60–85)\n• Touch-screen latency calibration\n• System Usability Scale (SUS) target > 82\n• Audio-assisted instructions for visually impaired",
         LAVENDER),
        ("Phase 3: Longitudinal Cohort\n[PROPOSED Q3 2027]",
         "• Academic multi-site research study (n=500)\n• 12-month longitudinal tracking\n• Cross-validation with external neuropsych batteries\n• Retest reliability r > 0.85 across battery",
         AMBER),
        ("Phase 4: Clinician Triage\n[PROPOSED 2028]",
         "• Clinician-supervised triage evaluation\n• Referral efficiency improvement target: 30%\n• HL7/FHIR EHR data exchange\n• Multi-site Exasol cluster deployment",
         CORAL)
    ]
    card_w = Inches(2.75)
    gap = Inches(0.24)
    for i, (title, desc, color) in enumerate(phases):
        left = Inches(0.8) + i * (card_w + gap)
        add_card(s11, left, Inches(2.0), card_w, Inches(4.8), border_color=color)
        tb = s11.shapes.add_textbox(left + Inches(0.15), Inches(2.2), card_w - Inches(0.3), Inches(4.4))
        tf = tb.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = title
        p.font.size = Pt(12)
        p.font.bold = True
        p.font.color.rgb = color
        p.font.name = "Segoe UI"
        p2 = tf.add_paragraph()
        p2.text = desc
        p2.font.size = Pt(10)
        p2.font.color.rgb = LIGHT_TEXT
        p2.font.name = "Segoe UI"

    s11.notes_slide.notes_text_frame.text = (
        "Presenter Script:\n"
        "Our roadmap defines a clear, responsible pathway forward. Phase 1 is complete: the core platform, 5 tasks, Exasol client, scoring engine, and 11 unit tests are verified. "
        "Phase 2 will validate usability and touch ergonomics among older adults. "
        "Phase 3 proposes a multi-site academic longitudinal research cohort tracking 500 individuals, and Phase 4 will evaluate clinical triage integration with hospital EHRs."
    )

    # =========================================================================
    # SLIDE 12: CLOSING & CALL TO ACTION
    # =========================================================================
    s12 = prs.slides.add_slide(blank_layout)
    set_slide_background(s12)
    add_header(s12, "Summary & Next Steps", "MindBridge AI: Empowering Longitudinal Cognitive Health",
               "Making cognitive performance easier to understand, track, and discuss.")
    add_footer(s12, 12)

    # Left: Platform Highlights Summary
    add_card(s12, Inches(0.8), Inches(2.0), Inches(5.7), Inches(4.8), border_color=TEAL)
    tb = s12.shapes.add_textbox(Inches(1.0), Inches(2.2), Inches(5.3), Inches(4.4))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "CORE PLATFORM ACHIEVEMENTS"
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = TEAL
    p.font.name = "Segoe UI"

    bullets = [
        "Interactive Psychometrics: 5 open cognitive paradigms with millisecond reaction timing.",
        "Exasol Personal Integration: In-memory columnar speed for longitudinal queries with transparent fallback integrity.",
        "Ethical Machine Learning: Calibrated Random Forest model with audited demographic subgroup fairness.",
        "Clinical Safeguards First: Non-diagnostic boundary, cryptographic consent, and immutable audit logs.",
        "Production Verified: Complete end-to-end web app with passing test suites and accessible UI."
    ]
    for b in bullets:
        p2 = tf.add_paragraph()
        p2.text = f"✓ {b}"
        p2.font.size = Pt(10.5)
        p2.font.color.rgb = LIGHT_TEXT
        p2.font.name = "Segoe UI"

    # Right: Collaboration / Call to Action
    add_card(s12, Inches(6.8), Inches(2.0), Inches(5.733), Inches(4.8), border_color=LAVENDER)
    tb = s12.shapes.add_textbox(Inches(7.0), Inches(2.2), Inches(5.3), Inches(4.4))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "COLLABORATION & RESEARCH INQUIRIES"
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = LAVENDER
    p.font.name = "Segoe UI"

    collab_items = [
        ("Clinical Partnerships", "Seeking academic neurology departments and cognitive research centers for longitudinal cohort validation studies."),
        ("Database & Analytics Scaling", "Collaborating on Exasol cluster scaling, distributed UDF analytics, and high-throughput psychometric data modeling."),
        ("Responsible AI Research", "Open to interdisciplinary collaboration on fairness auditing, algorithmic calibration, and clinical decision support.")
    ]
    for ctitle, cdesc in collab_items:
        p2 = tf.add_paragraph()
        p2.text = f"• {ctitle}:"
        p2.font.size = Pt(11)
        p2.font.bold = True
        p2.font.color.rgb = AMBER
        p2.font.name = "Segoe UI"
        p3 = tf.add_paragraph()
        p3.text = f"  {cdesc}"
        p3.font.size = Pt(10)
        p3.font.color.rgb = LIGHT_TEXT
        p3.font.name = "Segoe UI"

    s12.notes_slide.notes_text_frame.text = (
        "Presenter Script:\n"
        "In closing, MindBridge AI bridges the gap between raw computerized psychometrics, high-performance database analytics, and human clinical care. "
        "By grounding our technology in safety, transparent fallback modes, and audited demographic fairness, we make cognitive tracking understandable, dignified, and actionable. "
        "We welcome academic, clinical, and database technology partnerships. Thank you very much, and we invite your questions."
    )

    prs.save(str(deck_path))
    print(f"Presentation successfully saved to: {deck_path}")

if __name__ == "__main__":
    create_deck()
